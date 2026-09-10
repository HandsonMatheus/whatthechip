# -*- coding: utf-8 -*-
"""
A VOLTA DA PLANILHA: ler o arquivo que o comprador devolveu mexido.

O caso que motivou (EMIN-SO-2026-0004, Wu Quan, 09/2026): 105 categorias,
30 marcadas por ele com um check, e "ficou difícil de eu olhar e fazer um
resumo para entender exatamente no que ele mexeu". A auditoria achou 28
preços alterados — e TRÊS deles SEM check, um valendo ¥ 1.544 sozinho.

O que estes testes travam, e por que cada um existe:

  · O IDA-E-VOLTA. Exportar pelo caminho REAL (`compra_em_planilha`) e ler de
    volta tem de dar ZERO diferença. É a única trava que envelhece bem: no dia
    em que alguém renomear uma coluna do export, o parser para de casar e este
    teste fica vermelho — em vez de o comando relatar "105 linhas sumiram"
    contra uma planilha perfeitamente correta.
  · A MARCA HERDADA. O export ATUAL não tem coluna de marca (ela vive na
    faixa); o arquivo que ele devolveu, de uma versão anterior, tem. As duas
    formas precisam ler igual, senão o formato novo sai com 100% de "sumiu".
  · A DECOMPOSIÇÃO DO DINHEIRO. recusa + repreciação = a diferença inteira.
    Se as duas contas não somarem o total, o relatório está atribuindo a uma
    o que é da outra — e é essa atribuição que se leva para a negociação.
  · O ZERO SILENCIOSO. Ordem sem linha ABORTA. Sob RLS um comando fora de
    request lê zero em silêncio (CLAUDE.md §7), e aqui o zero se disfarça da
    melhor notícia possível: "nenhuma diferença encontrada".
"""

import io
from decimal import Decimal as D
from datetime import date

from django.core.management.base import CommandError
from django.contrib.auth import get_user_model
from django.test import TestCase

from estoque.models import Lot
from pricing.models import Buyer
from tenancy.models import Company, Membership
from tenancy.scope import company_scope
from vendas import planilha_auditoria as pa
from vendas.management.commands.auditar_planilha_comprador import (
    compara, linhas_da_ordem, referencia_da_aba_chips)
from vendas.models import (DocSequence, SEQ_SO, SalesOrder, SalesOrderLine,
                           STATUS_CONFIRMED)

User = get_user_model()
ZERO = D('0.00')


def _acha(linhas, marca, capacidade):
    """A linha da planilha por MARCA e CAPACIDADE. Endereçar por índice amarra
    o teste à ordenação do ORM (`kind, brand, gen, tier`) — e um teste que mexe
    na categoria errada passa a medir outra coisa sem avisar.

    ⚠ Não por CAIXA: sem `CategoryCode` semeado o código sai '—' em todas, e é
      assim mesmo — o que também exercita, de graça, a chave SECUNDÁRIA do
      `_chaves` (marca+tipo+capacidade), que é a que segura o caso real de
      categoria sem código de caixa."""
    for l in linhas:
        if l[0] == marca and l[2] == capacidade:
            return l
    raise AssertionError('sem linha %s/%s em %r' % (marca, capacidade, linhas))


def _wb(rows, titulo='EMIN-SO-2026-0004 · Wu Quan · LOT-2026-0007'):
    """Uma planilha no formato ANTIGO (com coluna Marca), em espanhol — a
    forma exata do arquivo que ele devolveu."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Resumen'
    ws.cell(1, 1, titulo)
    for i, t in enumerate(['Marca', 'Tipo', 'Capacidad', 'Caja WTC',
                           'Enviados', '¥ unit.', '¥ esperado', 'Rechazados',
                           'Aprobados', '¥ resultado'], start=1):
        ws.cell(3, i, t)
    # As larguras do arquivo REAL: A-H e J declaradas, I não. É a única forma
    # de a trava do deslocamento poder falhar — `insert_cols` move as células
    # e deixa as larguras onde estavam.
    for letra, w in (('A', 16.0), ('B', 14.0), ('C', 18.0), ('D', 14.0),
                     ('E', 11.0), ('F', 13.0), ('G', 15.0), ('H', 12.0),
                     ('J', 15.0)):
        ws.column_dimensions[letra].width = w
    for r, linha in enumerate(rows, start=4):
        for c, v in enumerate(linha, start=1):
            if v is not None:
                ws.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class NumeroTests(TestCase):
    """`numero()` — o que a planilha devolve não é o que o Excel escreveu."""

    def test_le_as_tres_grafias_de_milhar(self):
        self.assertEqual(pa.numero('1.234,56'), D('1234.56'))   # pt/es
        self.assertEqual(pa.numero('1,234.56'), D('1234.56'))   # en
        self.assertEqual(pa.numero('1234.56'), D('1234.56'))
        self.assertEqual(pa.numero('¥ 12'), D('12'))
        self.assertEqual(pa.numero(1234), D('1234'))

    def test_nao_inventa_zero(self):
        """'—' e vazio são AUSÊNCIA de preço. Virar 0 diria que ele zerou o
        valor — que é uma acusação, e uma que a planilha não faz."""
        for v in (None, '', '  ', '—', '-', 'abc', 'None'):
            self.assertIsNone(pa.numero(v), v)

    def test_booleano_nao_e_numero(self):
        """`True` é `int` em Python e viraria ¥ 1,00 — um preço inventado a
        partir de uma célula de caixa marcada."""
        self.assertIsNone(pa.numero(True))


class CabecalhoTests(TestCase):
    """O parser casa por NOME. Testado nas grafias que existem em produção."""

    def _papel(self, txt):
        return pa.papel_da_coluna(txt)

    def test_as_duas_grafias_do_unitario_sao_a_mesma_coluna(self):
        """'¥ unit.' (export até 09/2026) e 'UNITÁRIO ¥' (atual) são a MESMA
        coluna escrita por duas versões — e o arquivo que volta pode ser de
        qualquer uma delas."""
        self.assertEqual(self._papel('¥ unit.'), ('unit', False))
        self.assertEqual(self._papel('UNITÁRIO ¥'), ('unit', False))
        self.assertEqual(self._papel('UNITARIO ¥'), ('unit', False))

    def test_o_yen_separa_as_DUAS_colunas_de_recusa(self):
        """RECUSADOS (peças) e RECUSADOS ¥ (dinheiro) coexistem desde
        07/09/2026. Confundi-las põe dinheiro na coluna de quantidade."""
        self.assertEqual(self._papel('RECUSADOS'), ('recusados', False))
        self.assertEqual(self._papel('RECUSADOS ¥'), ('recusados', True))

    def test_coluna_que_o_sistema_nunca_escreveu_nao_tem_papel(self):
        for t in ('', None, '√', 'obs do comprador'):
            self.assertEqual(self._papel(t)[0], None, t)


class LeituraDoArquivoTests(TestCase):

    def test_faixa_de_marca_e_rodape_nao_sao_categoria(self):
        """Somar faixa junto com categoria dobra o lote inteiro — e o erro
        sai bonito: dá exatamente 2×."""
        buf = _wb([
            ['Micron', '', '', '', 10, '', 40, 0, 10, 40],          # faixa
            ['Micron', 'DDR3', '2Gb', 'E-08', 10, 4, 40, 0, 10, 40],
            ['Total · 1 marcas', '', '', '', 10, '', 40, 0, 10, 40],
        ])
        import openpyxl
        r = pa.le_resumo(openpyxl.load_workbook(buf).active)
        self.assertEqual(len(r['linhas']), 1)
        self.assertEqual(len(r['faixas']), 1)
        self.assertEqual(r['total']['enviados'], D('10'))

    def test_check_do_comprador_em_coluna_que_o_sistema_nao_escreveu(self):
        buf = _wb([
            ['Micron', 'DDR3', '2Gb', 'E-08', 10, 4, 40, 0, 10, 40, None, '√'],
            ['Micron', 'DDR3', '4Gb', 'E-03', 5, 4, 20, 0, 5, 20, None, ' '],
            ['Micron', 'DDR4', '8Gb', 'E-06', 5, 4, 20, 0, 5, 20],
        ])
        import openpyxl
        ls = pa.le_resumo(openpyxl.load_workbook(buf).active)['linhas']
        self.assertEqual([l['check'] for l in ls], [True, False, False])
        # ⚠ A célula com UM ESPAÇO existe no arquivo real (Samsung LPDDR3 4GB).
        #   Espaço não é marca — mas também não é célula intocada: alguém
        #   digitou ali. Tratar as duas como a mesma coisa esconde a única
        #   linha em que a intenção dele é ambígua.
        self.assertEqual([l['check_ilegivel'] for l in ls],
                         [False, True, False])

    def test_marca_e_HERDADA_quando_o_export_nao_tem_a_coluna(self):
        """O export ATUAL tirou a coluna Marca (a faixa já diz a marca). Sem a
        herança, toda linha do formato novo sai com marca vazia e NENHUMA casa
        com o banco — 100% de 'sumiu', que é o relatório mais alarmante e mais
        falso possível."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        for i, t in enumerate(['TIPO', 'CAPACIDADE', 'CAIXA WTC',
                               'UNITÁRIO ¥', 'ENVIADOS', 'ESPERADO ¥'],
                              start=1):
            ws.cell(2, i, t)
        # A faixa MESCLA TIPO..WTC e escreve 'Micron   1 linha' em TIPO.
        ws.cell(3, 1, 'Micron   1 linha')
        ws.cell(3, 5, 10)
        for i, v in enumerate(['DDR3', '2Gb', 'E-08', 4, 10, 40], start=1):
            ws.cell(4, i, v)
        r = pa.le_resumo(ws)
        self.assertEqual(len(r['linhas']), 1)
        self.assertEqual(r['linhas'][0]['marca'], 'Micron')


class _ComOrdem(TestCase):
    """Uma OV confirmada de 3 categorias, com o ¥ CONGELADO na linha."""

    CATEGORIAS = (
        # (marca, kind, gen, tier, unidade, qtd, ¥ congelado)
        ('Samsung', 'emmc', '', D('8'), 'GB', 216, D('20.00')),
        ('Samsung', 'emcp', 'LPDDR4', D('16'), 'GB', 355, D('17.00')),
        ('Micron', 'ddr', 'DDR3', D('2'), 'Gb', 601, D('3.00')),
    )

    def setUp(self):
        self.emp = Company.objects.create(name='eMiner', slug='eminer',
                                          code='EMIN')
        self.buyer = Buyer.all_companies.create(company=None, name='Wu Quan',
                                                slug='wu-quan')
        self.user = User.objects.create_user('g', password='x')
        Membership.objects.create(user=self.user, company=self.emp,
                                  role=Membership.ROLE_MANAGER)
        with company_scope(self.emp.id):
            self.lot = Lot.all_companies.create(
                company=self.emp, number=7, description='x', status='closed',
                operator=self.user, origin='pcb')
            self.so = SalesOrder(
                lot=self.lot, buyer=self.buyer, status=STATUS_CONFIRMED,
                fx_usd_rate=D('0.1400'), total_rmb=D('9999.00'),
                total_usd=D('1399.86'), shipped_at=date(2026, 8, 27),
                number=DocSequence.next_number(self.emp, SEQ_SO))
            self.so.save()
            for marca, kind, gen, tier, un, qtd, unit in self.CATEGORIAS:
                SalesOrderLine.all_companies.create(
                    order=self.so, company=self.emp, brand=marca, kind=kind,
                    gen=gen, tier_value=tier, tier_unit=un, quantity=qtd,
                    unit_rmb=unit, unit_usd=(unit * D('0.14')).quantize(
                        D('0.01')))

    def _banco(self):
        return linhas_da_ordem(self.so.code, None)[1]

    def _corpo(self, mexer=None, com_faixa=False):
        """As LINHAS da planilha como ela saiu, com o que `mexer` pedir.

        `mexer` recebe a lista (marca, tipo, cap, caixa, env, unit, esperado,
        recusa, aprov, resultado, _, check) e a devolve mexida — é o comprador
        editando. `com_faixa` põe a faixa da marca em cima, como no arquivo
        real (o export agrupa por marca)."""
        linhas = []
        for b in self._banco():
            env, unit = b['qty'], b['unit']
            linhas.append([b['marca'], b['tipo'], b['capacidade'], b['caixa'],
                           env, float(unit), float(unit * env), 0, env,
                           float(unit * env), None, None])
        if mexer:
            linhas = mexer(linhas)
        if com_faixa:
            soma = sum(l[9] or 0 for l in linhas)
            por_marca, com = {}, []
            for l in linhas:
                por_marca.setdefault(l[0], []).append(l)
            for marca, suas in por_marca.items():
                com.append([marca, '', '', '', 0, None, 0, 0, 0, 0,
                            None, None])
                com += suas
            linhas = com + [['Total · %d marcas' % len(por_marca), '', '', '',
                             0, None, soma, 0, 0, soma, None, None]]
        return linhas

    def _planilha(self, mexer=None):
        import openpyxl
        return pa.le_resumo(
            openpyxl.load_workbook(_wb(self._corpo(mexer))).active)


class ComparaTests(_ComOrdem):

    def test_planilha_intacta_nao_acha_diferenca(self):
        """A âncora: sem esta, um comparador que acha diferença em TUDO
        passaria em todos os outros testes."""
        achados = compara(self._planilha(), self._banco())
        self.assertEqual(len(achados), 3)
        self.assertTrue(all(a['tipo_achado'] == 'linha' and not a['diffs']
                            for a in achados), [a['diffs'] for a in achados])

    def test_preco_alterado_com_o_de_para_e_o_impacto(self):
        def baixa_o_emmc(linhas):
            l = _acha(linhas, 'Samsung', '8GB')
            l[5] = 15.0                              # ¥ 20 → 15
            l[9] = 216 * 15.0
            l[11] = '√'
            return linhas
        achados = compara(self._planilha(baixa_o_emmc), self._banco())
        mexida = [a for a in achados if a['diffs']]
        self.assertEqual(len(mexida), 1)
        a = mexida[0]
        self.assertEqual(a['diffs'], ['preco'])
        self.assertEqual(a['banco']['unit'], D('20.00'))
        self.assertEqual(a['planilha']['unit'], D('15'))
        self.assertEqual(a['perda_preco'], D('1080'))     # 216 × 5
        self.assertTrue(a['planilha']['check'])

    def test_recusa_nao_e_alteracao_de_preco(self):
        """As duas contas não podem se misturar: recusa é direito dele;
        repreciação é proposta. Somadas num número só, a diferença 'fecha' —
        e é por fechar que passar batido é fácil."""
        def recusa_30(linhas):
            l = _acha(linhas, 'Samsung', '16GB')
            l[7], l[8], l[9] = 30, 355 - 30, (355 - 30) * 17.0
            return linhas
        a = [x for x in compara(self._planilha(recusa_30), self._banco())
             if x['diffs']][0]
        self.assertEqual(a['diffs'], ['recusa'])
        self.assertEqual(a['perda_recusa'], D('510'))     # 30 × 17
        self.assertEqual(a['perda_preco'], ZERO)

    def test_mexer_na_quantidade_ENVIADA_e_achado_proprio(self):
        """ENVIADOS é o que o cliente despachou — ele informa recusa na coluna
        de recusa. Reescrever o enviado apaga a referência do despacho."""
        def some_9(linhas):
            _acha(linhas, 'Micron', '2Gb')[4] = 601 - 9
            return linhas
        a = [x for x in compara(self._planilha(some_9), self._banco())
             if x['diffs']][0]
        self.assertIn('enviados', a['diffs'])

    def test_a_conta_dele_que_nao_fecha_vira_achado(self):
        """No arquivo real há linha com `aprovados` que não é
        `enviados − recusados` e `¥ resultado` que não é `aprovados × ¥`.
        Nenhuma das duas é diferença de PREÇO — e some se só o preço for
        conferido."""
        def desalinha(linhas):
            l = _acha(linhas, 'Samsung', '8GB')
            l[7], l[8] = 10, 216       # recusa 10 e aprova 216 mesmo assim
            return linhas
        a = [x for x in compara(self._planilha(desalinha), self._banco())
             if x['diffs']][0]
        self.assertIn('conta_aprovados', a['diffs'])

    def test_linha_que_sumiu_e_linha_que_apareceu(self):
        def bagunca(linhas):
            linhas.remove(_acha(linhas, 'Samsung', '8GB'))
            linhas.append(['Kingston', 'eMMC', '64GB', 'B-07', 5, 30.0,
                           150.0, 0, 5, 150.0, None, None])
            return linhas
        achados = compara(self._planilha(bagunca), self._banco())
        tipos = sorted(a['tipo_achado'] for a in achados)
        self.assertEqual(tipos, ['linha', 'linha', 'so_na_planilha',
                                 'so_no_banco'])

    def test_a_decomposicao_do_dinheiro_FECHA(self):
        """recusa + repreciação = a diferença inteira, linha a linha. Se as
        duas não somarem o total, o relatório atribui a uma o que é da outra —
        e é essa atribuição que se leva para a negociação."""
        def as_duas_coisas(linhas):
            l = _acha(linhas, 'Samsung', '8GB')
            l[5] = 15.0                    # ¥ 20 → 15
            l[7], l[8] = 16, 200           # e recusa 16
            l[9] = 200 * 15.0
            return linhas
        a = [x for x in compara(self._planilha(as_duas_coisas), self._banco())
             if x['diffs']][0]
        self.assertEqual(a['valor_ov'], D('4320'))        # 216 × 20
        self.assertEqual(a['valor_dele'], D('3000'))      # 200 × 15
        self.assertEqual(a['perda_recusa'], D('320'))     # 16 × 20
        self.assertEqual(a['perda_preco'], D('1000'))     # 200 × 5
        self.assertEqual(a['valor_ov'] - a['perda_recusa'] - a['perda_preco'],
                         a['valor_dele'])


class IdaEVoltaTests(_ComOrdem):
    """A trava que envelhece bem: exportar pelo caminho REAL e ler de volta.

    Todos os outros testes desta suíte constroem a planilha à mão — e uma
    planilha construída à mão prova o parser contra a MINHA ideia do formato,
    não contra o formato. No dia em que alguém renomear uma coluna do
    `planilha.py`, só este teste fica vermelho; sem ele, o comando passaria a
    relatar '105 categorias sumiram' contra um arquivo perfeitamente correto,
    e o relatório mais assustador possível seria o mais errado.
    """

    def _exporta(self):
        from vendas.planilha import compra_em_planilha
        from vendas.views_partner import _detalhe
        import openpyxl
        # ⚠ DENTRO do escopo: `_detalhe` toca InventoryEntry, que é
        #   multi-empresa (mesma ressalva de tests_planilha_da_compra).
        with company_scope(self.emp.id):
            dados, _nome = compra_em_planilha(self.so, _detalhe(self.so))
        return openpyxl.load_workbook(io.BytesIO(dados), data_only=True)

    def test_o_arquivo_que_o_sistema_ACABOU_de_emitir_nao_tem_diferenca(self):
        wb = self._exporta()
        resumo = pa.le_resumo(wb.worksheets[0])
        self.assertEqual(len(resumo['linhas']), 3, resumo['papeis'])
        achados = compara(resumo, self._banco())
        self.assertEqual([a['tipo_achado'] for a in achados],
                         ['linha'] * 3)
        self.assertEqual([a['diffs'] for a in achados], [[], [], []])

    def test_a_faixa_da_marca_do_export_real_nao_vira_categoria(self):
        """A faixa MESCLA TIPO..WTC e escreve 'Samsung   2 linhas' na célula
        de TIPO — ou seja, ela TEM tipo preenchido. A primeira versão deste
        parser reconhecia faixa por 'tipo vazio' e lia as duas faixas como
        categorias: o lote saía com 5 linhas em vez de 3, e a soma dobrava as
        marcas com mais de uma categoria."""
        resumo = pa.le_resumo(self._exporta().worksheets[0])
        self.assertEqual(len(resumo['faixas']), 2)          # Micron, Samsung
        self.assertEqual(sorted(f['marca'] for f in resumo['faixas']),
                         ['Micron', 'Samsung'])
        self.assertNotIn('linha', [f['marca'] for f in resumo['faixas']])

    def test_um_preco_mexido_no_arquivo_real_aparece_e_so_ele(self):
        """Morde a mutação óbvia: um comparador que devolvesse [] sempre
        passaria no teste do ida-e-volta limpo."""
        wb = self._exporta()
        ws = wb.worksheets[0]
        col = pa.mapa_de_colunas(ws)[1]
        alvo = [l for l in pa.le_resumo(ws)['linhas']
                if l['capacidade'] == '8GB'][0]
        ws.cell(alvo['linha'], col['unit'], 12)             # ¥ 20 → 12
        achados = compara(pa.le_resumo(ws), self._banco())
        mexidas = [a for a in achados if a['diffs']]
        self.assertEqual(len(mexidas), 1)
        self.assertEqual(mexidas[0]['diffs'], ['preco'])
        self.assertEqual(mexidas[0]['perda_preco'], D('1728'))   # 216 × 8


class ZeroSilenciosoTests(_ComOrdem):
    """Ordem sem linha ABORTA — nunca vira 'nenhuma diferença encontrada'.

    Sob RLS, comando fora de request lê ZERO em silêncio (CLAUDE.md §7). Este
    comando é o pior lugar do sistema para isso acontecer: o zero se disfarça
    da melhor notícia possível, e a melhor notícia possível é o que autoriza
    pagar sem discutir. Em SQLite não há RLS — o teste não simula a policy,
    trava a REAÇÃO ao conjunto vazio, que é o que precisa existir."""

    def test_ordem_sem_linha_levanta_em_vez_de_relatar_tudo_certo(self):
        with company_scope(self.emp.id):
            SalesOrderLine.all_companies.filter(order=self.so).delete()
        with self.assertRaises(CommandError) as e:
            linhas_da_ordem(self.so.code, None)
        self.assertIn('ZERO linhas', str(e.exception))

    def test_codigo_inexistente_levanta_falando_do_DATABASE_URL(self):
        """Apontar para o banco errado é a outra forma silenciosa de este
        comando não achar diferença nenhuma."""
        with self.assertRaises(CommandError) as e:
            linhas_da_ordem('EMIN-SO-1999-9999', None)
        self.assertIn('DATABASE_URL', str(e.exception))

    def test_le_a_ordem_certa_e_o_yuan_CONGELADO(self):
        _ov, linhas = linhas_da_ordem(self.so.code, None)
        self.assertEqual(len(linhas), 3)
        self.assertEqual(
            sorted((l['marca'], l['tipo'], l['capacidade'], l['qty'],
                    l['unit']) for l in linhas),
            [('Micron', 'DDR3', '2Gb', 601, D('3.00')),
             ('Samsung', 'LPDDR4', '16GB', 355, D('17.00')),
             ('Samsung', 'eMMC', '8GB', 216, D('20.00'))])


class TestemunhaDaAbaChipsTests(TestCase):
    """`--sem-banco`: a aba Chips reconstrói o preço que SAIU do sistema."""

    def test_agrega_por_marca_e_caixa(self):
        chips = {'linhas': [
            {'pn': 'A1', 'marca': 'Micron', 'caixa': 'E-08', 'qtd': D('2'),
             'unit': D('3')},
            {'pn': 'A2', 'marca': 'Micron', 'caixa': 'E-08', 'qtd': D('4'),
             'unit': D('3')},
        ]}
        ref, conflitos = referencia_da_aba_chips(chips)
        self.assertEqual(ref[('micron', 'e-08')]['qty'], D('6'))
        self.assertEqual(ref[('micron', 'e-08')]['unit'], D('3'))
        self.assertEqual(conflitos, [])

    def test_testemunha_contaminada_e_DENUNCIADA_nao_escolhida(self):
        """Se ele editou a aba Chips também, dois PNs do mesmo par trazem ¥
        diferente. Escolher um dos dois em silêncio faria a auditoria inteira
        repousar sobre um palpite — e ela seria apresentada como prova."""
        chips = {'linhas': [
            {'pn': 'A1', 'marca': 'Micron', 'caixa': 'E-08', 'qtd': D('2'),
             'unit': D('3')},
            {'pn': 'A2', 'marca': 'Micron', 'caixa': 'E-08', 'qtd': D('4'),
             'unit': D('2')},
        ]}
        _ref, conflitos = referencia_da_aba_chips(chips)
        self.assertEqual(len(conflitos), 1)
        self.assertEqual(conflitos[0][4], 'A2')


class PintaAPlanilhaDeleTests(_ComOrdem):
    """O entregável é a planilha DELE, marcada — e as colunas novas ENTRAM NO
    MEIO, encostadas no preço.

    Três rodadas com o dono neste mesmo arquivo, e as três falhas foram de
    LUGAR e NOME, nunca de conta:

      1. comparativo de 17 colunas em arquivo novo → *"impossível ler,
         informação demais"*;
      2. as colunas certas chamadas `¥ SISTEMA`/`¥ DELE` → ele pediu *"uma
         coluna com o preço COMO ERA ANTES"*, que era exatamente aquela;
      3. acrescentadas DEPOIS da última coluna dele → *"não veio coluna de
         antes e agora não"*. Vieram — estavam a três colunas de rolagem.

    Daí o que esta classe trava: informação a uma tela de distância do dado
    que ela explica não existe, e rótulo tem de responder à pergunta de quem
    lê ("quanto era?"), não nomear a fonte do dado ("o sistema").
    """

    def _pinta(self, mexer):
        """Escreve o arquivo NO DISCO e parseia ELE — nunca uma segunda cópia.

        ⚠ Os achados carregam NÚMERO DE LINHA. Parsear um workbook e pintar
          outro (ainda que "igual") desloca tudo em uma linha assim que os dois
          deixarem de ser idênticos — foi o que aconteceu ao acrescentar a
          faixa da marca a só um dos dois.
        """
        import openpyxl
        import tempfile
        import os
        ent = os.path.join(tempfile.mkdtemp(), 'dele.xlsx')
        # ⚠ COM FAIXA DE MARCA E RODAPÉ, como o arquivo real: sem faixa, a
        #   linha acima da 1ª categoria é o próprio cabeçalho e a dedução
        #   errada acerta por acaso; sem rodapé, o total não tem onde somar.
        with open(ent, 'wb') as f:
            f.write(_wb(self._corpo(mexer, com_faixa=True)).read())
        resumo = pa.le_resumo(openpyxl.load_workbook(ent).active)
        achados = compara(resumo, self._banco())
        sai = ent.replace('.xlsx', ' - PINTADA.xlsx')
        from vendas.planilha_auditoria_saida import pinta_planilha
        contas = pinta_planilha(ent, sai, achados, resumo)
        return openpyxl.load_workbook(sai).worksheets[0], contas, ent, resumo

    def _baixa(self, linhas):
        l = _acha(linhas, 'Samsung', '8GB')
        l[5], l[9] = 15.0, 216 * 15.0            # ¥ 20 → 15
        return linhas

    def _cats(self, ws):
        """As linhas de categoria da planilha pintada (têm caixa WTC)."""
        return [r for r in range(4, ws.max_row + 1) if ws.cell(r, 4).value]

    # ── o LUGAR das colunas, que foi a queixa ───────────────────────────
    def test_as_colunas_novas_entram_ENCOSTADAS_no_preco(self):
        """`¥ ANTES` imediatamente à ESQUERDA do ¥ unitário dele, e os dois
        totais imediatamente à direita.

        É a trava da 3ª rodada. Acrescentadas no fim (que é o que `append`
        faria), elas caem depois das colunas que o comprador inventou e sob a
        borda direita da tela — e uma coluna que só aparece rolando não foi
        entregue.
        """
        ws, _c, _e, _r = self._pinta(self._baixa)
        titulos = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        i = titulos.index('¥ ANTES')
        self.assertEqual(titulos[i:i + 4],
                         ['¥ ANTES', '¥ unit.', '¥ TOTAL ANTES',
                          '¥ TOTAL AGORA'])

    def test_as_colunas_dele_continuam_todas_la_e_na_ordem(self):
        """Inserir no meio empurra o resto para a direita — nada pode ser
        sobrescrito nem trocar de ordem. A planilha tem de continuar a dele."""
        ws, _c, _e, _r = self._pinta(self._baixa)
        titulos = [t for t in (ws.cell(3, c).value
                               for c in range(1, ws.max_column + 1)) if t]
        dele = [t for t in titulos if t not in
                ('¥ ANTES', '¥ TOTAL ANTES', '¥ TOTAL AGORA')]
        self.assertEqual(dele, ['Marca', 'Tipo', 'Capacidad', 'Caja WTC',
                                'Enviados', '¥ unit.', '¥ esperado',
                                'Rechazados', 'Aprobados', '¥ resultado'])

    def test_as_LARGURAS_acompanham_o_deslocamento(self):
        """`insert_cols` move as células mas NÃO as larguras — elas são um
        dicionário à parte, indexado por LETRA. Sem remontá-las, tudo à direita
        do preço herda a largura da coluna vizinha, e coluna estreita demais
        mostra `####` no lugar do número.

        A prova é a largura de `¥ resultado`: 15,0 na coluna J do arquivo de
        entrada, e tem de continuar 15,0 depois de a coluna virar M. E o
        descritor do openpyxl recusa `width = None`, então "limpar" o holder
        errado também estoura aqui.
        """
        ws, _c, _e, _r = self._pinta(self._baixa)
        titulos = {ws.cell(3, c).value: c
                   for c in range(1, ws.max_column + 1)}
        self.assertGreater(titulos['¥ resultado'], 10)   # J → M
        largura = lambda t: ws.column_dimensions[
            ws.cell(3, titulos[t]).column_letter].width
        self.assertEqual(largura('¥ resultado'), 15.0)
        self.assertEqual(largura('Marca'), 16.0)         # A não se move
        self.assertEqual(largura('¥ unit.'), 13.0)       # F → G, com a dele
        self.assertEqual(largura('¥ esperado'), 15.0)    # G → J

    def test_o_titulo_novo_cai_na_LINHA_DE_CABECALHO_dele(self):
        """A 1ª versão deduzia a linha do cabeçalho como 'a primeira linha de
        dado menos um'. No arquivo real, a linha acima da 1ª categoria é a
        FAIXA DA MARCA — e o título ia parar dentro dela, invisível."""
        ws, _c, _e, _r = self._pinta(self._baixa)
        self.assertIn('¥ ANTES',
                      [ws.cell(3, c).value
                       for c in range(1, ws.max_column + 1)])

    # ── o que é pintado ─────────────────────────────────────────────────
    def test_a_celula_do_PRECO_dele_e_que_fica_pintada(self):
        ws, contas, _e, _r = self._pinta(self._baixa)
        titulos = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        c_antes = titulos.index('¥ ANTES') + 1
        alvo = [r for r in self._cats(ws) if ws.cell(r, 3).value == '8GB'][0]
        self.assertTrue(ws.cell(alvo, c_antes + 1).fill.patternType)
        self.assertEqual([ws.cell(alvo, c).value
                          for c in (c_antes, c_antes + 1)], [20.0, 15.0])
        self.assertEqual((contas['baixou'], contas['subiu']), (1, 0))

    def test_subir_e_baixar_nao_levam_a_mesma_cor(self):
        def sobe(linhas):
            l = _acha(linhas, 'Micron', '2Gb')
            l[5], l[9] = 4.0, 601 * 4.0          # ¥ 3 → 4
            return linhas

        def cor(ws, cap):
            titulos = [ws.cell(3, c).value
                       for c in range(1, ws.max_column + 1)]
            c = titulos.index('¥ ANTES') + 2
            r = [r for r in self._cats(ws) if ws.cell(r, 3).value == cap][0]
            return ws.cell(r, c).fill.fgColor.rgb
        ws_b, cb, _e, _r = self._pinta(self._baixa)
        ws_s, cs, _e2, _r2 = self._pinta(sobe)
        self.assertNotEqual(cor(ws_b, '8GB'), cor(ws_s, '2Gb'))
        self.assertEqual((cb['baixou'], cs['subiu']), (1, 1))

    def test_recusa_sozinha_NAO_e_marcada(self):
        """A coluna de recusados é dele, é onde ele deve escrever. Marcá-la
        trataria o combinado — 'mandei 10, chegou 1 quebrado, pago 9' — como
        irregularidade. O que se marca é o que ele mudou onde não devia."""
        def recusa(linhas):
            l = _acha(linhas, 'Samsung', '8GB')
            l[7], l[8], l[9] = 16, 200, 200 * 20.0
            return linhas
        ws, contas, _e, _r = self._pinta(recusa)
        titulos = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        c_agora = titulos.index('¥ ANTES') + 2
        alvo = [r for r in self._cats(ws) if ws.cell(r, 3).value == '8GB'][0]
        self.assertFalse(ws.cell(alvo, c_agora).fill.patternType)
        self.assertEqual((contas['baixou'], contas['subiu']), (0, 0))

    # ── as duas colunas de total ────────────────────────────────────────
    def test_o_total_da_linha_e_ENVIADOS_x_antes_e_APROVADOS_x_dele(self):
        """As duas contas são diferentes de propósito. ANTES é o lote como foi
        fechado: **enviados** × ¥ congelado. AGORA é o que ele propõe pagar:
        **aprovados** × ¥ dele. Usar enviados nas duas esconderia a recusa;
        usar aprovados nas duas esconderia que a recusa custou algo."""
        def as_duas_coisas(linhas):
            l = _acha(linhas, 'Samsung', '8GB')
            l[5] = 15.0                    # ¥ 20 → 15
            l[7], l[8] = 16, 200           # e recusa 16 de 216
            l[9] = 200 * 15.0
            return linhas
        ws, contas, _e, _r = self._pinta(as_duas_coisas)
        titulos = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        c_ta = titulos.index('¥ TOTAL ANTES') + 1
        alvo = [r for r in self._cats(ws) if ws.cell(r, 3).value == '8GB'][0]
        self.assertEqual(ws.cell(alvo, c_ta).value, 216 * 20.0)
        self.assertEqual(ws.cell(alvo, c_ta + 1).value, 200 * 15.0)
        # 216×20 + 355×17 + 601×3 = 12158
        self.assertEqual(contas['antes'], D('12158'))
        self.assertEqual(contas['agora'], D('10838'))

    def test_a_faixa_soma_as_linhas_DELA_e_o_rodape_soma_as_FAIXAS(self):
        """A conta que não pode dobrar. Faixa que soma o intervalo errado, ou
        rodapé que soma as linhas outra vez em cima das faixas, dá exatamente
        2× — e o erro sai bonito, num número redondo que ninguém questiona."""
        import re
        ws, _c, _e, resumo = self._pinta(self._baixa)
        titulos = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        c = titulos.index('¥ TOTAL ANTES') + 1
        cats = set(self._cats(ws))
        faixas = [f['linha'] for f in resumo['faixas']]
        self.assertGreater(len(faixas), 1, 'cenário precisa de 2+ faixas')

        coberto = []
        for b in faixas:
            m = re.match(r'=SUM\([A-Z]+(\d+):[A-Z]+(\d+)\)$',
                         str(ws.cell(b, c).value))
            self.assertIsNotNone(m, 'faixa %d sem SUM: %r'
                                 % (b, ws.cell(b, c).value))
            coberto += list(range(int(m.group(1)), int(m.group(2)) + 1))
        # cada categoria somada UMA vez, e nenhuma faixa dentro do intervalo
        self.assertEqual(sorted(coberto), sorted(cats))
        self.assertEqual(len(coberto), len(set(coberto)))

        rodape = str(ws.cell(resumo['total']['linha'], c).value)
        self.assertEqual(sorted(int(n) for n in re.findall(r'\d+', rodape)),
                         sorted(faixas))

    def test_ANTES_existe_em_TODA_linha_inclusive_nas_intocadas(self):
        """Dono, 08/09/2026: *"uma coluna com o preço COMO ERA ANTES, pra eu
        ter ideia do que baixou"*. Coluna com buraco não se corre o olho nem
        se soma — e varrer a planilha de cima a baixo é a comparação que ele
        quer fazer. O que separa mexido de intacto é a COR, não a ausência.

        ⚠ Esta trava nasceu de uma mutação que PASSOU: os outros testes desta
          classe olham linhas COM diferença, e a versão que só preenchia essas
          passava em todos eles.
        """
        ws, _c, _e, _r = self._pinta(self._baixa)
        titulos = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        c_antes = titulos.index('¥ ANTES') + 1
        intactas = 0
        for r in self._cats(ws):
            self.assertIsNotNone(ws.cell(r, c_antes).value,
                                 'linha %d sem ¥ ANTES' % r)
            if not ws.cell(r, c_antes + 1).fill.patternType:
                self.assertEqual(ws.cell(r, c_antes).value,
                                 ws.cell(r, c_antes + 1).value, 'linha %d' % r)
                intactas += 1
        self.assertEqual(intactas, 2)            # 3 categorias, 1 mexida

    def test_o_arquivo_que_ele_mandou_nao_e_tocado(self):
        """O original é a prova do que ele mandou, e é a prova que sustenta a
        conversa. A pintura vai sempre para uma cópia."""
        import hashlib
        _ws, _c, entrada, _r = self._pinta(self._baixa)
        antes = hashlib.sha256(open(entrada, 'rb').read()).hexdigest()
        self._pinta(self._baixa)
        self.assertEqual(hashlib.sha256(open(entrada, 'rb').read()).hexdigest(),
                         antes)
